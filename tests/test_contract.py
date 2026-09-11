"""Cross-source contract checks.

Every source flattens its own record dataclass into ``dict[str, str]`` rows keyed
by the headers in its ``columns`` module, and ``common.excel.write_workbook``
silently drops any key that is not a declared header.  So a column/attribute
mismatch introduced by a refactor would not fail a parser test — it would just
surface as a blank column in the workbook.  These tests pin the contract:

* ``COLUMNS`` maps real dataclass attributes to unique headers, and
  ``HEADERS`` is exactly those headers in order;
* ``rows_from_records`` emits rows whose keys are exactly ``HEADERS``;
* the end-to-end workbook's header row is exactly ``HEADERS``.
"""

from __future__ import annotations

import dataclasses
from collections.abc import Callable
from typing import Any, NamedTuple

import pytest
from openpyxl import load_workbook

from sanctions_lists_etl import run_source
from sanctions_lists_etl.sources.eu import columns as eu_columns
from sanctions_lists_etl.sources.eu import parser as eu_parser
from sanctions_lists_etl.sources.ofac import columns as ofac_columns
from sanctions_lists_etl.sources.ofac import parser as ofac_parser
from sanctions_lists_etl.sources.uk import columns as uk_columns
from sanctions_lists_etl.sources.uk import parser as uk_parser
from sanctions_lists_etl.sources.un import columns as un_columns
from sanctions_lists_etl.sources.un import parser as un_parser


class Contract(NamedTuple):
    name: str
    record_cls: type
    columns: list[tuple[str, str]]
    headers: list[str]
    parse: Callable[[Any], list]
    rows_from_records: Callable[[list], list[dict[str, str]]]
    fixture: str  # conftest fixture name for the sample file
    run_kwarg: str  # run_source kwarg that points at a local file
    sheet: str


CONTRACTS: list[Contract] = [
    Contract(
        "ofac",
        ofac_parser.PartyRecord,
        ofac_columns.COLUMNS,
        ofac_columns.HEADERS,
        ofac_parser.parse_sdn_advanced,
        ofac_parser.rows_from_records,
        "sample_xml",
        "xml_path",
        "SDN",
    ),
    Contract(
        "eu",
        eu_parser.SanctionEntity,
        eu_columns.COLUMNS,
        eu_columns.HEADERS,
        eu_parser.parse_eu_fsf,
        eu_parser.rows_from_records,
        "sample_eu_xml",
        "xml_path",
        "EU",
    ),
    Contract(
        "un",
        un_parser.SanctionParty,
        un_columns.COLUMNS,
        un_columns.HEADERS,
        un_parser.parse_un_consolidated,
        un_parser.rows_from_records,
        "sample_un_xml",
        "xml_path",
        "UN",
    ),
    Contract(
        "uk",
        uk_parser.SanctionParty,
        uk_columns.COLUMNS,
        uk_columns.HEADERS,
        uk_parser.parse_uk_sanctions,
        uk_parser.rows_from_records,
        "sample_uk_xml",
        "xml_path",
        "UK",
    ),
]

_IDS = [c.name for c in CONTRACTS]


@pytest.mark.parametrize("contract", CONTRACTS, ids=_IDS)
def test_columns_map_real_attributes_to_unique_headers(contract: Contract):
    attrs = {f.name for f in dataclasses.fields(contract.record_cls)}
    col_attrs = [attr for attr, _ in contract.columns]
    col_headers = [header for _, header in contract.columns]

    unknown = [a for a in col_attrs if a not in attrs]
    assert not unknown, f"{contract.name}: COLUMNS references unknown attributes {unknown}"
    assert len(col_headers) == len(set(col_headers)), f"{contract.name}: duplicate headers"
    assert contract.headers == col_headers, f"{contract.name}: HEADERS out of sync with COLUMNS"


@pytest.mark.parametrize("contract", CONTRACTS, ids=_IDS)
def test_rows_have_exactly_the_declared_headers(contract: Contract, request):
    sample = request.getfixturevalue(contract.fixture)
    rows = contract.rows_from_records(contract.parse(sample))
    assert rows, f"{contract.name}: fixture produced no rows"
    want = set(contract.headers)
    for row in rows:
        assert set(row) == want, f"{contract.name}: row keys {set(row) ^ want} differ from headers"


@pytest.mark.parametrize("contract", CONTRACTS, ids=_IDS)
def test_workbook_header_row_matches(contract: Contract, request, tmp_path):
    sample = request.getfixturevalue(contract.fixture)
    result = run_source(
        contract.name,
        output_dir=tmp_path,
        raw_dir=tmp_path,
        download=False,
        **{contract.run_kwarg: sample},
    )
    sheet = load_workbook(result.xlsx_path)[contract.sheet]
    header_row = [cell.value for cell in sheet[1]]
    assert header_row == contract.headers
    assert sheet.max_row == result.record_count + 1
