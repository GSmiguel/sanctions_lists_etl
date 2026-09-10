"""The OFAC Consolidated (Non-SDN) list rides on the same parser and columns as
the SDN list; only the download URL, workbook name and sheet differ."""

import shutil

import pytest
from openpyxl import load_workbook

from sanctions_lists_etl.cli import build_parser
from sanctions_lists_etl.runner import run_source
from sanctions_lists_etl.sources.ofac.pipeline import options_from_args


def _cache_as(sample_xml, raw_dir, name):
    dest = raw_dir / name
    shutil.copy(sample_xml, dest)
    return dest


def test_consolidated_run_writes_its_own_workbook(sample_xml, tmp_path):
    _cache_as(sample_xml, tmp_path, "cons_advanced.xml")

    result = run_source(
        "ofac",
        output_dir=tmp_path,
        raw_dir=tmp_path,
        lists=("consolidated",),
        download=False,
    )

    dest = tmp_path / "ofac_consolidated.xlsx"
    assert dest.exists()
    assert not (tmp_path / "ofac_sdn.xlsx").exists()
    assert load_workbook(dest).sheetnames[0] == "CONS"
    assert result.record_count == 4
    assert result.metadata["lists_built"] == "consolidated"


def test_both_lists_produce_both_workbooks(sample_xml, tmp_path):
    _cache_as(sample_xml, tmp_path, "sdn_advanced.xml")
    _cache_as(sample_xml, tmp_path, "cons_advanced.xml")

    result = run_source(
        "ofac",
        output_dir=tmp_path,
        raw_dir=tmp_path,
        lists=("sdn", "consolidated"),
        download=False,
    )

    assert (tmp_path / "ofac_sdn.xlsx").exists()
    assert (tmp_path / "ofac_consolidated.xlsx").exists()
    assert result.record_count == 8  # 4 + 4
    assert result.metadata["sdn_record_count"] == "4"
    assert result.metadata["consolidated_record_count"] == "4"


@pytest.mark.parametrize(
    "argv, expected",
    [
        (["ofac"], ("sdn", "consolidated")),
        (["ofac", "--list", "sdn"], ("sdn",)),
        (["ofac", "--list", "consolidated"], ("consolidated",)),
        (["ofac", "--list", "both"], ("sdn", "consolidated")),
    ],
)
def test_list_flag_maps_to_lists_option(argv, expected):
    args = build_parser().parse_args(argv)
    assert options_from_args(args)["lists"] == expected
