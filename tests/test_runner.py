import pytest
from openpyxl import load_workbook

from sanctions_lists_etl import available_sources, run_source, runner
from sanctions_lists_etl.base import Source, SourceResult
from sanctions_lists_etl.cli import main


def test_sources_are_registered():
    assert {"ofac", "eu", "un", "uk"} <= set(available_sources())


def test_run_source_end_to_end(sample_xml, tmp_path):
    result = run_source(
        "ofac",
        output_dir=tmp_path,
        raw_dir=tmp_path,
        xml_path=sample_xml,
        download=False,
    )
    assert result.source == "ofac"
    assert result.record_count == 4
    assert result.xlsx_path.exists()
    assert load_workbook(result.xlsx_path)["SDN"].max_row == 5


def test_run_eu_source_end_to_end(sample_eu_xml, tmp_path):
    result = run_source(
        "eu",
        output_dir=tmp_path,
        raw_dir=tmp_path,
        xml_path=sample_eu_xml,
        download=False,
    )
    assert result.source == "eu"
    assert result.record_count == 5
    assert result.xlsx_path.exists()
    assert load_workbook(result.xlsx_path)["EU"].max_row == 6


def test_run_un_source_end_to_end(sample_un_xml, tmp_path):
    result = run_source(
        "un",
        output_dir=tmp_path,
        raw_dir=tmp_path,
        xml_path=sample_un_xml,
        download=False,
    )
    assert result.source == "un"
    assert result.record_count == 7
    assert result.counts_by_type == {"Individual": 4, "Entity": 3}
    assert result.xlsx_path.exists()
    assert load_workbook(result.xlsx_path)["UN"].max_row == 8


def test_run_uk_source_end_to_end(sample_uk_xml, tmp_path):
    result = run_source(
        "uk",
        output_dir=tmp_path,
        raw_dir=tmp_path,
        xml_path=sample_uk_xml,
        download=False,
    )
    assert result.source == "uk"
    assert result.record_count == 5
    assert result.counts_by_type == {"Entity": 2, "Individual": 2, "Vessel": 1}
    assert result.xlsx_path.exists()
    assert load_workbook(result.xlsx_path)["UK"].max_row == 6


def test_cli_eu_requires_token_without_xml(tmp_path, capsys, monkeypatch):
    for name in ("EU_FSF_TOKEN", "EU_FSF_TOKEN_FILE", "EU_FSF_URL"):
        monkeypatch.delenv(name, raising=False)
    exit_code = main(
        ["--output-dir", str(tmp_path), "--raw-dir", str(tmp_path), "eu", "--no-download"]
    )
    assert exit_code == 1
    assert "EU_FSF_TOKEN" in capsys.readouterr().err


def test_cli_runs_single_source(sample_xml, tmp_path, capsys):
    exit_code = main(
        [
            "--output-dir",
            str(tmp_path),
            "--raw-dir",
            str(tmp_path),
            "ofac",
            "--xml",
            str(sample_xml),
            "--individuals-entities-only",
        ]
    )
    assert exit_code == 0
    assert (tmp_path / "ofac_sdn.xlsx").exists()
    out = capsys.readouterr().out
    assert "[ofac]" in out
    assert "Vessel" not in out


def test_cli_list(capsys):
    assert main(["--list"]) == 0
    assert "ofac" in capsys.readouterr().out


def test_run_all_exclude_skips_named_sources(monkeypatch, tmp_path):
    ran: list[str] = []

    def make_run(name):
        def _run(*, output_dir, raw_dir, **_):
            ran.append(name)
            return SourceResult(source=name, xlsx_path=tmp_path / f"{name}.xlsx", record_count=0)

        return _run

    monkeypatch.setattr(
        runner,
        "_SOURCES",
        {n: Source(name=n, description="", run=make_run(n)) for n in ("a", "b", "c")},
    )

    runner.run_all(output_dir=tmp_path, raw_dir=tmp_path, exclude=["b"])
    assert ran == ["a", "c"]


def test_run_all_exclude_rejects_unknown_source(monkeypatch, tmp_path):
    monkeypatch.setattr(
        runner, "_SOURCES", {"a": Source(name="a", description="", run=lambda **_: None)}
    )
    with pytest.raises(KeyError, match="nope"):
        runner.run_all(output_dir=tmp_path, raw_dir=tmp_path, exclude=["nope"])


def test_cli_exclude_only_with_all(sample_xml, capsys):
    with pytest.raises(SystemExit) as exc:
        main(["--exclude", "eu", "ofac", "--xml", str(sample_xml)])
    assert exc.value.code == 2
    assert "--exclude only applies" in capsys.readouterr().err


def test_run_all_keeps_going_past_a_failing_source(monkeypatch, tmp_path):
    def ok_run(*, output_dir, raw_dir, **_):
        return SourceResult(source="ok", xlsx_path=tmp_path / "ok.xlsx", record_count=1)

    def boom_run(*, output_dir, raw_dir, **_):
        raise RuntimeError("no credential")

    monkeypatch.setattr(
        runner,
        "_SOURCES",
        {
            "ok": Source(name="ok", description="", run=ok_run),
            "boom": Source(name="boom", description="", run=boom_run),
        },
    )

    with pytest.raises(RuntimeError, match="boom: no credential"):
        runner.run_all(output_dir=tmp_path, raw_dir=tmp_path)
