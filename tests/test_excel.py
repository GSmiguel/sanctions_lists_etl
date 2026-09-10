from openpyxl import load_workbook

from sanctions_lists_etl.common.excel import write_workbook
from sanctions_lists_etl.sources.ofac.columns import HEADERS
from sanctions_lists_etl.sources.ofac.parser import parse_sdn_advanced, rows_from_records


def test_write_workbook_roundtrip(sample_xml, tmp_path):
    rows = rows_from_records(parse_sdn_advanced(sample_xml))
    dest = write_workbook(
        HEADERS, rows, tmp_path / "out.xlsx", sheet_name="SDN", metadata={"source": "test"}
    )

    workbook = load_workbook(dest)
    sheet = workbook["SDN"]
    assert [cell.value for cell in sheet[1]] == HEADERS
    assert sheet.max_row == len(rows) + 1
    assert sheet.freeze_panes == "A2"
    assert "info" in workbook.sheetnames


def test_long_values_are_clipped(tmp_path):
    rows = [{header: ("x" * 40_000 if header == "addresses" else "") for header in HEADERS}]
    dest = write_workbook(HEADERS, rows, tmp_path / "big.xlsx")
    sheet = load_workbook(dest)["data"]
    assert len(sheet.cell(row=2, column=HEADERS.index("addresses") + 1).value) <= 32_000
